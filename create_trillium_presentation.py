from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUT = "/home/shivam/ideas/trillium-technology-project-ideas.pptx"

NAVY = RGBColor(16, 42, 67)
BLUE = RGBColor(30, 94, 145)
TEAL = RGBColor(0, 151, 157)
GREEN = RGBColor(83, 166, 92)
ORANGE = RGBColor(239, 139, 44)
INK = RGBColor(32, 45, 58)
MUTED = RGBColor(91, 107, 121)
LIGHT = RGBColor(242, 246, 248)
PALE_BLUE = RGBColor(229, 240, 247)
PALE_TEAL = RGBColor(226, 245, 244)
PALE_ORANGE = RGBColor(254, 241, 225)
WHITE = RGBColor(255, 255, 255)
LINE = RGBColor(208, 220, 226)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def add_rect(slide, x, y, w, h, fill, radius=False, line=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line if line else fill
    return shape


def add_text(slide, text, x, y, w, h, size=20, color=INK, bold=False,
             align=PP_ALIGN.LEFT, font="Aptos", valign=MSO_ANCHOR.TOP,
             margin=0.04):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_rich_text(slide, parts, x, y, w, h, size=18, color=INK,
                  align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    for text, bold, part_color in parts:
        run = p.add_run()
        run.text = text
        run.font.name = "Aptos"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = part_color or color
    return box


def add_title(slide, title, subtitle=None, section=None):
    if section:
        add_text(slide, section.upper(), 0.62, 0.30, 3.0, 0.28, 10, TEAL, True)
    add_text(slide, title, 0.62, 0.63, 12.0, 0.55, 27, NAVY, True)
    add_rect(slide, 0.62, 1.25, 0.72, 0.055, TEAL)
    if subtitle:
        add_text(slide, subtitle, 0.62, 1.39, 11.9, 0.48, 13.5, MUTED)


def add_footer(slide, number, source=None):
    add_rect(slide, 0.62, 7.10, 12.05, 0.012, LINE)
    if source:
        add_text(slide, source, 0.65, 7.16, 11.1, 0.18, 8, MUTED)
    add_text(slide, str(number), 12.18, 7.13, 0.42, 0.20, 9, MUTED, True,
             align=PP_ALIGN.RIGHT)


def add_bullet_list(slide, items, x, y, w, h, size=16, color=INK,
                    bullet_color=TEAL, gap=0.55):
    for i, item in enumerate(items):
        cy = y + i * gap
        add_rect(slide, x, cy + 0.10, 0.10, 0.10, bullet_color, radius=True)
        add_text(slide, item, x + 0.24, cy, w - 0.24, gap, size, color)


def add_card(slide, x, y, w, h, title, body, accent=TEAL, number=None):
    add_rect(slide, x, y, w, h, WHITE, radius=True, line=LINE)
    add_rect(slide, x, y, 0.09, h, accent, radius=True)
    if number is not None:
        add_rect(slide, x + 0.30, y + 0.25, 0.54, 0.54, accent, radius=True)
        add_text(slide, str(number), x + 0.30, y + 0.25, 0.54, 0.54, 17,
                 WHITE, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        title_x = x + 1.00
        title_w = w - 1.25
    else:
        title_x = x + 0.35
        title_w = w - 0.65
    add_text(slide, title, title_x, y + 0.24, title_w, 0.40, 18, NAVY, True)
    add_text(slide, body, x + 0.35, y + 0.85, w - 0.65, h - 1.05, 13.5, MUTED)


def add_arrow(slide, x, y, w=0.52, h=0.28, color=TEAL):
    shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = color
    return shape


def add_table(slide, rows, col_widths, x, y, h, header_fill=NAVY):
    table_shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y),
                                         Inches(sum(col_widths)), Inches(h))
    table = table_shape.table
    for i, width in enumerate(col_widths):
        table.columns[i].width = Inches(width)
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = value
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.05)
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if r == 0 else (WHITE if r % 2 else LIGHT)
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT
                for run in p.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(11.5 if r else 11)
                    run.font.bold = r == 0 or c == 0
                    run.font.color.rgb = WHITE if r == 0 else INK
    return table_shape


# Slide 1 — Title
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_rect(slide, 0, 0, 13.333, 7.5, NAVY)
add_rect(slide, 0, 0, 0.18, 7.5, TEAL)
add_rect(slide, 9.83, 0, 3.50, 7.5, BLUE)
add_rect(slide, 10.30, 1.02, 2.22, 2.22, TEAL, radius=True)
add_rect(slide, 10.77, 1.49, 1.28, 1.28, NAVY, radius=True)
add_rect(slide, 10.30, 4.24, 2.22, 0.20, ORANGE, radius=True)
add_text(slide, "TECHNOLOGY PROJECT PROPOSAL", 0.75, 1.00, 6.8, 0.35, 12, TEAL, True)
add_text(slide, "Two opportunities to advance\nTrillium’s pump business", 0.75, 1.58, 8.4, 1.55,
         31, WHITE, True)
add_text(slide, "From AI-assisted engineering to closed-loop customer service",
         0.77, 3.46, 7.9, 0.54, 17, RGBColor(206, 223, 235))
add_rect(slide, 0.77, 5.50, 2.12, 0.46, TEAL, radius=True)
add_text(slide, "CONCEPT DISCUSSION", 0.77, 5.50, 2.12, 0.46, 10, WHITE, True,
         align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
add_text(slide, "September 2026", 0.77, 6.23, 2.4, 0.30, 11, RGBColor(190, 211, 224))

# Slide 2 — Why now
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, "Why this is the right time", "Three business signals create a practical opening for software-led improvement.", "Context")
add_card(slide, 0.72, 2.10, 3.75, 3.72, "A sharper pump focus",
         "Recent portfolio changes increase the importance of differentiated pump products, services, and aftermarket growth.", BLUE, 1)
add_card(slide, 4.79, 2.10, 3.75, 3.72, "AI research is underway",
         "Trillium and Politecnico di Milano are developing fast surrogate models for centrifugal-pump simulation and geometry.", TEAL, 2)
add_card(slide, 8.86, 2.10, 3.75, 3.72, "Competitors are digitizing",
         "Major competitors connect pump selection and monitoring with recommendations, workflows, APIs, and service offerings.", ORANGE, 3)
add_text(slide, "The opportunity: convert existing engineering and monitoring strengths into usable digital products.",
         1.25, 6.32, 10.85, 0.45, 17, NAVY, True, align=PP_ALIGN.CENTER)
add_footer(slide, 2, "Sources: Trillium Flow Technologies; AIRIC / Politecnico di Milano; competitor public product pages")

# Slide 3 — Opportunity map
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, "Two complementary opportunities", "One improves how customers maintain pumps; the other improves how Trillium designs them.", "Opportunity")
add_rect(slide, 0.78, 2.05, 5.65, 3.85, PALE_TEAL, radius=True, line=PALE_TEAL)
add_text(slide, "CUSTOMER-FACING", 1.15, 2.38, 2.5, 0.28, 10.5, TEAL, True)
add_text(slide, "PumpCare 360", 1.15, 2.83, 4.8, 0.46, 24, NAVY, True)
add_text(slide, "Turn pump alerts into clear maintenance and service actions.", 1.15, 3.45, 4.55, 0.72, 16, INK)
add_bullet_list(slide, ["Lower downtime", "Improve customer experience", "Grow digital and aftermarket revenue"],
                1.18, 4.38, 4.75, 1.42, 14, MUTED, TEAL, 0.45)
add_rect(slide, 6.90, 2.05, 5.65, 3.85, PALE_BLUE, radius=True, line=PALE_BLUE)
add_text(slide, "INTERNAL", 7.27, 2.38, 2.5, 0.28, 10.5, BLUE, True)
add_text(slide, "RapidDesign", 7.27, 2.83, 4.8, 0.46, 24, NAVY, True)
add_text(slide, "Turn AI pump research into a faster design-to-quote workflow.", 7.27, 3.45, 4.55, 0.72, 16, INK)
add_bullet_list(slide, ["Reduce engineering cycle time", "Evaluate more design options", "Respond to customers faster"],
                7.30, 4.38, 4.75, 1.42, 14, MUTED, BLUE, 0.45)
add_footer(slide, 3)

# Slide 4 — PumpCare overview
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, "Idea 1 — PumpCare 360", "Extend Cyclops² from condition visibility to a complete action and service experience.", "Customer platform")
add_text(slide, "TODAY", 0.78, 2.05, 1.1, 0.26, 10, MUTED, True)
add_rect(slide, 0.78, 2.42, 3.32, 2.75, LIGHT, radius=True, line=LINE)
add_text(slide, "Monitor", 1.08, 2.77, 2.7, 0.38, 20, NAVY, True)
add_text(slide, "Cyclops² provides connected pump data, condition monitoring, and analytics.",
         1.08, 3.43, 2.70, 1.10, 14, MUTED)
add_arrow(slide, 4.38, 3.57, 0.72, 0.40, TEAL)
add_text(slide, "PROPOSED", 5.32, 2.05, 1.4, 0.26, 10, TEAL, True)
add_rect(slide, 5.32, 2.42, 7.20, 2.75, PALE_TEAL, radius=True, line=PALE_TEAL)
for i, (title, body) in enumerate([
    ("Explain", "What changed and how serious is it?"),
    ("Recommend", "What action should the customer take?"),
    ("Execute", "Create a task, service request, or parts enquiry."),
]):
    x = 5.66 + i * 2.25
    add_rect(slide, x, 2.82, 1.88, 0.52, TEAL if i < 2 else GREEN, radius=True)
    add_text(slide, title, x, 2.82, 1.88, 0.52, 14, WHITE, True,
             align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, body, x, 3.62, 1.88, 0.92, 12.5, INK, align=PP_ALIGN.CENTER)
add_rich_text(slide, [("Business outcome: ", True, NAVY),
                      ("lower downtime + stronger service relationship + recurring revenue", False, INK)],
              1.22, 5.80, 10.90, 0.48, 17, align=PP_ALIGN.CENTER)
add_footer(slide, 4, "Source: Trillium Smart Condition Monitoring / Cyclops² public materials")

# Slide 5 — PumpCare journey
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, "PumpCare 360 — how value is created", "A closed loop connects operating data to customer action and Trillium services.", "Concept flow")
flow = [
    ("1", "Pump data", "Cyclops², sensors, or approved customer data", BLUE),
    ("2", "Prioritized insight", "Health, anomaly, urgency, and explanation", TEAL),
    ("3", "Recommended action", "Engineer-approved maintenance guidance", GREEN),
    ("4", "Execution", "Work order, service request, or parts enquiry", ORANGE),
]
for i, (num, title, body, color) in enumerate(flow):
    x = 0.72 + i * 3.14
    add_rect(slide, x, 2.35, 2.68, 3.10, WHITE, radius=True, line=LINE)
    add_rect(slide, x + 0.22, 2.60, 0.56, 0.56, color, radius=True)
    add_text(slide, num, x + 0.22, 2.60, 0.56, 0.56, 17, WHITE, True,
             align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, title, x + 0.25, 3.42, 2.18, 0.55, 17, NAVY, True, align=PP_ALIGN.CENTER)
    add_text(slide, body, x + 0.30, 4.16, 2.08, 0.83, 12.7, MUTED, align=PP_ALIGN.CENTER)
    if i < 3:
        add_arrow(slide, x + 2.76, 3.66, 0.38, 0.24, TEAL)
add_rect(slide, 1.35, 5.95, 10.62, 0.60, NAVY, radius=True)
add_text(slide, "Every resolved issue improves the asset history and supports future recommendations.",
         1.35, 5.95, 10.62, 0.60, 15, WHITE, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
add_footer(slide, 5)

# Slide 6 — Monitoring competition
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, "Competitor signal: monitoring is becoming actionable", "The public market benchmark extends beyond dashboards and alerts.", "Market comparison")
rows = [
    ["Competitor", "Public offering", "Visible strength"],
    ["KSB", "KSB Guard", "Maintenance tracking, action recommendations, expert support, REST API"],
    ["Sulzer", "BLUE BOX", "Risk insights plus energy, cost, carbon, and remaining-life analysis"],
    ["Flowserve", "RedRaven", "Connected predictive monitoring across flow-control equipment"],
    ["Xylem", "Avensor", "Anomaly detection, prioritization, and guided inspections"],
]
add_table(slide, rows, [1.55, 2.10, 7.75], 0.82, 2.08, 3.32)
add_rect(slide, 0.82, 5.77, 11.40, 0.82, PALE_ORANGE, radius=True, line=PALE_ORANGE)
add_rich_text(slide, [("Trillium opportunity: ", True, NAVY),
                      ("connect Cyclops² insight directly to maintenance execution and Trillium aftermarket services.", False, INK)],
              1.12, 5.95, 10.80, 0.42, 15.5, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
add_footer(slide, 6, "Sources: KSB Guard, Sulzer BLUE BOX, Flowserve RedRaven, Xylem Avensor public pages")

# Slide 7 — RapidDesign overview
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, "Idea 2 — RapidDesign", "Productize existing AI simulation research into an everyday engineering and quotation tool.", "Internal platform")
add_rect(slide, 0.78, 2.10, 4.02, 3.92, LIGHT, radius=True, line=LINE)
add_text(slide, "Current challenge", 1.10, 2.46, 3.34, 0.42, 19, NAVY, True)
add_bullet_list(slide, [
    "Custom pump designs require multiple iterations",
    "Traditional CFD is computationally expensive",
    "Knowledge is distributed across experts and past projects",
    "Slow technical response can weaken competitiveness",
], 1.12, 3.16, 3.24, 2.45, 13.2, MUTED, ORANGE, 0.57)
add_arrow(slide, 5.05, 3.68, 0.72, 0.40, TEAL)
add_rect(slide, 6.02, 2.10, 6.48, 3.92, PALE_BLUE, radius=True, line=PALE_BLUE)
add_text(slide, "RapidDesign response", 6.38, 2.46, 5.72, 0.42, 19, NAVY, True)
add_bullet_list(slide, [
    "Capture customer operating requirements",
    "Retrieve relevant approved designs and curves",
    "Use AI models to compare alternatives quickly",
    "Send the best candidates for final CFD validation",
    "Generate a controlled technical proposal",
], 6.40, 3.12, 5.48, 2.68, 13.2, INK, BLUE, 0.50)
add_footer(slide, 7, "Source: Trillium / AIRIC Reduced Order Modelling and PIAI4PUMPS-GEO project information")

# Slide 8 — RapidDesign flow and competition
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, "RapidDesign — faster without removing engineering control", "AI narrows the search space; qualified engineers retain final design authority.", "Concept flow")
steps = [
    ("Requirements", "Duty point, fluid, standards, constraints"),
    ("AI screening", "Evaluate and rank many design scenarios"),
    ("Engineering validation", "Review shortlist and run final CFD"),
    ("Proposal", "Approved curves, assumptions, and configuration"),
]
for i, (title, body) in enumerate(steps):
    x = 0.69 + i * 3.16
    add_rect(slide, x, 2.13, 2.75, 2.15, WHITE, radius=True, line=LINE)
    add_rect(slide, x, 2.13, 2.75, 0.11, BLUE if i != 1 else TEAL)
    add_text(slide, title, x + 0.18, 2.56, 2.39, 0.44, 16, NAVY, True, align=PP_ALIGN.CENTER)
    add_text(slide, body, x + 0.25, 3.25, 2.25, 0.65, 12.3, MUTED, align=PP_ALIGN.CENTER)
    if i < 3:
        add_arrow(slide, x + 2.79, 3.02, 0.36, 0.22, TEAL)
add_text(slide, "COMPETITOR POSITIONING", 0.78, 4.78, 2.8, 0.25, 10, TEAL, True)
add_rect(slide, 0.78, 5.15, 11.78, 1.27, PALE_TEAL, radius=True, line=PALE_TEAL)
add_text(slide, "Flowserve Affinity, KSB EasySelect, and Sulzer Select accelerate catalogue selection.",
         1.05, 5.40, 5.45, 0.58, 14, INK)
add_text(slide, "RapidDesign differentiates through AI-assisted custom engineering—not just standard product selection.",
         6.72, 5.34, 5.48, 0.70, 14, NAVY, True)
add_footer(slide, 8, "Sources: Flowserve Affinity, KSB EasySelect, Sulzer Select / ABSEL public pages")

# Slide 9 — Comparison
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, "How the two ideas compare", "Both are valuable, but their dependencies and time-to-value are different.", "Decision")
rows = [
    ["Decision area", "PumpCare 360", "RapidDesign"],
    ["Primary users", "Customers and service teams", "Engineering and sales teams"],
    ["Core value", "Turn alerts into maintenance and service", "Reduce design and quotation time"],
    ["Revenue path", "Subscriptions, parts, and services", "Faster response and more competitive bids"],
    ["Main dependency", "Cyclops² data and customer integrations", "Engineering data and existing AI models"],
    ["Recommended timing", "Strategic second initiative", "Start first"],
]
add_table(slide, rows, [2.25, 4.55, 4.55], 0.94, 2.12, 3.92)
add_rect(slide, 1.90, 6.27, 9.55, 0.55, NAVY, radius=True)
add_text(slide, "Recommendation: begin with RapidDesign; develop PumpCare 360 as the customer-facing follow-on.",
         1.90, 6.27, 9.55, 0.55, 14.5, WHITE, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
add_footer(slide, 9)

# Slide 10 — Pilot
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, "Recommended first move: a focused RapidDesign pilot", "Validate business value before defining the detailed implementation roadmap.", "Next step")
add_card(slide, 0.78, 2.10, 3.72, 3.78, "1. Confirm the baseline",
         "Select one centrifugal-pump family and document current quotation time, CFD effort, revisions, and available data.", BLUE, 1)
add_card(slide, 4.81, 2.10, 3.72, 3.78, "2. Build a controlled pilot",
         "Connect requirements, approved design history, and the existing AI model in a simple engineer-facing workflow.", TEAL, 2)
add_card(slide, 8.84, 2.10, 3.72, 3.78, "3. Compare results",
         "Run historical or live cases and compare prediction accuracy, engineering time, alternatives evaluated, and proposal speed.", GREEN, 3)
add_text(slide, "Success gate: measurable cycle-time reduction with acceptable model accuracy and full engineer approval.",
         1.14, 6.30, 11.0, 0.44, 16, NAVY, True, align=PP_ALIGN.CENTER)
add_footer(slide, 10)

# Slide 11 — Decision and sources
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_rect(slide, 0, 0, 13.333, 7.5, NAVY)
add_text(slide, "DECISION REQUEST", 0.74, 0.60, 2.4, 0.30, 11, TEAL, True)
add_text(slide, "Approve a short discovery phase", 0.74, 1.03, 7.9, 0.62, 28, WHITE, True)
add_text(slide, "Bring together engineering, digital products, service, and sales to validate:",
         0.76, 1.90, 8.1, 0.42, 15, RGBColor(205, 222, 234))
add_bullet_list(slide, [
    "Existing tools and overlapping initiatives",
    "Available engineering models and datasets",
    "A suitable RapidDesign pump family and pilot owner",
    "The Cyclops² roadmap and future PumpCare 360 opportunity",
], 0.82, 2.62, 7.55, 2.62, 15, WHITE, TEAL, 0.60)
add_rect(slide, 9.20, 0.72, 3.33, 5.75, RGBColor(25, 57, 84), radius=True, line=RGBColor(25, 57, 84))
add_text(slide, "KEY PUBLIC SOURCES", 9.56, 1.08, 2.6, 0.32, 11, TEAL, True)
sources = [
    "Trillium Smart Condition Monitoring",
    "AIRIC Reduced Order Modelling",
    "Trillium Data-Centre Cooling Pumps",
    "KSB Guard and EasySelect",
    "Sulzer BLUE BOX and Select",
    "Flowserve RedRaven and Affinity",
    "Xylem Avensor",
]
add_bullet_list(slide, sources, 9.55, 1.70, 2.56, 3.58, 11.5,
                RGBColor(220, 232, 240), TEAL, 0.48)
add_text(slide, "Publicly visible gaps must be validated against Trillium’s internal roadmap.",
         9.56, 5.59, 2.58, 0.58, 10.5, RGBColor(179, 203, 218))
add_text(slide, "Detailed implementation planning follows team approval.",
         0.76, 6.48, 7.9, 0.42, 16, WHITE, True)
add_text(slide, "11", 12.20, 7.12, 0.40, 0.20, 9, RGBColor(180, 201, 215), True,
         align=PP_ALIGN.RIGHT)

prs.save(OUT)
print(OUT)
print(f"slides={len(prs.slides)}")
